"""Regenera la ficha técnica editable ``ficha_repositorio.docx``.

La ficha lee los CSV producidos por ``report/generate_report_assets.py`` para
que porcentajes, intervalos y benchmarks no se copien manualmente. Ejecución
recomendada desde ``zombie_poisson_streamlit``::

    python report/generate_report_assets.py
    python presentation/generate_ficha.py

El PDF homónimo se compila desde ``ficha_repositorio.tex``; se mantiene una
fuente LaTeX independiente porque convertir DOCX a PDF no es portable entre
Windows, Linux y los entornos de integración continua.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


PRESENTATION_DIR = Path(__file__).resolve().parent
PROJECT_DIR = PRESENTATION_DIR.parent
REPOSITORY_DIR = PROJECT_DIR.parent
ASSETS_DIR = PROJECT_DIR / "assets"
DATA_DIR = PROJECT_DIR / "report" / "data"
OUTPUT_PATH = PRESENTATION_DIR / "ficha_repositorio.docx"
REPOSITORY_URL = (
    "https://github.com/DanielBarillasM/"
    "Proyecto-1_Grupo-1_Modelacion-y-Simulacion_Sec-30"
)

DARK = "0B100D"
GREEN = "9FE870"
BLUE = "68B8D8"
MUTED = "607068"


def set_cell_shading(cell, colour: str) -> None:
    """Aplica un relleno OOXML; python-docx no expone esta opción directamente."""

    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), colour)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    """Inserta un hipervínculo externo activo en un párrafo de Word."""

    relationship = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    colour = OxmlElement("w:color")
    colour.set(qn("w:val"), "2E75B6")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend((colour, underline))
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.extend((properties, text_element))
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def style_document(document: Document) -> None:
    """Configura tipografía, márgenes, encabezado y pie de página."""

    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    for name, size, colour in (
        ("Title", 30, DARK),
        ("Heading 1", 19, DARK),
        ("Heading 2", 13, "276749"),
    ):
        style = document.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(colour)

    header = section.header.paragraphs[0]
    header.text = "OUTBREAK  /  STOCHASTIC SURVIVAL LAB"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    footer = section.footer.paragraphs[0]
    footer.text = "Proyecto 1 · Modelación y Simulación · Sección 30"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor.from_string(MUTED)


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    """Añade una tabla compacta con cabecera coherente con la identidad visual."""

    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Light Shading Accent 1"
    for cell, label in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, DARK)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        run = cell.paragraphs[0].add_run(label)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(GREEN)
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    document.add_paragraph()


def add_bullets(document: Document, items: list[str]) -> None:
    """Crea una lista breve; cada punto expresa una afirmación verificable."""

    for item in items:
        document.add_paragraph(item, style="List Bullet")


def load_evidence() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Carga evidencia generada y falla temprano si falta algún entregable."""

    names = (
        "comparison_summary.csv",
        "paired_effects.csv",
        "normal_method_tests.csv",
        "normal_method_decision.csv",
    )
    missing = [name for name in names if not (DATA_DIR / name).is_file()]
    if missing:
        raise FileNotFoundError(
            "Faltan resultados; ejecute report/generate_report_assets.py: "
            + ", ".join(missing)
        )
    return tuple(pd.read_csv(DATA_DIR / name) for name in names)  # type: ignore[return-value]


def main() -> None:
    """Construye una ficha autocontenida, editable y sincronizada con los CSV."""

    summary, effects, normal_tests, normal_decision = load_evidence()
    poisson = summary.set_index("model").loc["poisson"]
    polar = summary.set_index("model").loc["polar"]
    survival_effect = effects.set_index("metric").loc["survived"]
    normal_quality = normal_tests.set_index("method")
    normal_scores = normal_decision.set_index("method")

    document = Document()
    style_document(document)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("OUTBREAK\n")
    subtitle = title.add_run("Stochastic Survival Lab")
    subtitle.font.color.rgb = RGBColor.from_string("276749")
    document.add_picture(str(ASSETS_DIR / "outbreak-command-center.png"), width=Inches(6.8))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    lead = document.add_paragraph()
    lead.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lead.add_run("Ficha técnica del repositorio\n").bold = True
    lead.add_run("Grupo 1 · CC2017 · Sección 30 · 2026")
    link = document.add_paragraph()
    link.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_hyperlink(link, "Abrir repositorio en GitHub", REPOSITORY_URL)

    document.add_heading("1. Propósito y pregunta", level=1)
    document.add_paragraph(
        "Se modela el subsistema real de aparición y balance de enemigos de un "
        "videojuego. La simulación conecta tiempos aleatorios de llegada con "
        "movimiento, combate y probabilidad de completar una misión de 90 s."
    )
    callout = document.add_paragraph()
    callout.add_run("Pregunta central. ").bold = True
    callout.add_run(
        "¿Cómo cambian congestión y supervivencia al conservar la media, "
        "pero modificar la distribución de los interarribos?"
    )
    add_bullets(document, [
        "Poisson homogéneo: interarribos exponenciales por transformada inversa.",
        "Polar--lognormal: renovación positiva con media 1/lambda y CV configurable.",
        "Comparación algorítmica: Marsaglia Polar y Box--Muller generan la misma N(0,1).",
        "Monte Carlo pareado: Wilson, McNemar exacta, Wilcoxon y bootstrap.",
    ])

    document.add_heading("2. Modelo e implementación", level=1)
    add_table(document, ["Componente", "Implementación", "Propósito"], [
        ["Llegadas", "Delta = -ln(U)/lambda", "Poisson de tasa constante"],
        ["Contraste", "Delta = exp(mu + sigma Z)", "Media igual; variabilidad distinta"],
        ["Combate", "Paso fijo dt = 0.05 s", "Movimiento, alcance, contacto y daño"],
        ["Visual", "Plotly Mesh3d", "Escena navegable dentro de Streamlit"],
    ])
    add_table(document, ["Parámetro base", "Valor", "Unidad"], [
        ["Tasa", "55.8", "llegadas/min"],
        ["CV Polar", "0.60", "adimensional"],
        ["Horizonte", "90", "s"],
        ["HP / DPS protagonista", "110 / 42", "HP / HP por s"],
        ["HP / velocidad / DPS infectado", "42 / 1.55 / 9.5", "HP / m por s / HP por s"],
    ])
    document.add_paragraph(
        "La calibración evita un escenario saturado cerca de 100 %; no representa "
        "telemetría real y se acompaña de un barrido de tasas."
    )

    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_heading("3. Validación y comparación de métodos", level=1)
    document.add_paragraph(
        "La batería contiene 18 contrastes con corrección Holm: fuente uniforme "
        "(PCG64 y LCG propio), transformaciones, independencia, aceptación y "
        "conteo Poisson. Un contador uniforme pero predecible actúa como control "
        "negativo y falla las tres pruebas de dependencia."
    )
    document.add_picture(str(ASSETS_DIR / "report-normal-method-comparison.png"), width=Inches(6.7))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_table(document, ["Método normal", "KS p", "Rechazo", "Puntaje / 5"], [
        [
            "Marsaglia Polar",
            f"{normal_quality.loc['Marsaglia Polar', 'p_value']:.4f}",
            f"{normal_quality.loc['Marsaglia Polar', 'rejected_fraction'] * 100:.1f} %",
            f"{normal_scores.loc['Marsaglia Polar', 'weighted_total']:.3f}",
        ],
        [
            "Box--Muller",
            f"{normal_quality.loc['Box-Muller', 'p_value']:.4f}",
            "0 %",
            f"{normal_scores.loc['Box-Muller', 'weighted_total']:.3f}",
        ],
    ])
    document.add_paragraph(
        "Ambos métodos son compatibles con N(0,1). La matriz favorece a "
        f"{normal_decision.iloc[0]['method']} en el equipo de referencia; no "
        "interpreta un valor p mayor como mejor ajuste."
    )

    document.add_heading("4. Resultado reproducible", level=1)
    add_table(document, ["Métrica (400 corridas)", "Poisson", "Polar--lognormal"], [
        ["Supervivencia", f"{poisson.survival_probability * 100:.1f} %", f"{polar.survival_probability * 100:.1f} %"],
        ["IC Wilson 95 %", f"[{poisson.ci_low * 100:.1f}, {poisson.ci_high * 100:.1f}] %", f"[{polar.ci_low * 100:.1f}, {polar.ci_high * 100:.1f}] %"],
        ["Tiempo medio", f"{poisson.mean_survival_time:.3f} s", f"{polar.mean_survival_time:.3f} s"],
        ["Pico activo medio", f"{poisson.mean_max_concurrent:.3f}", f"{polar.mean_max_concurrent:.3f}"],
    ])
    document.add_paragraph(
        f"Efecto Poisson menos Polar: {survival_effect.difference * 100:.1f} "
        f"puntos porcentuales, IC [{survival_effect.ci_low * 100:.1f}, "
        f"{survival_effect.ci_high * 100:.1f}]. La conclusión es condicional "
        "al escenario; mayor supervivencia no equivale a mejor modelo."
    )

    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_heading("5. Arquitectura, uso y entregables", level=1)
    add_table(document, ["Ruta", "Responsabilidad"], [
        ["App/simulation.py", "Motor, generadores, validación e inferencia"],
        ["App/visuals.py", "Gráficas Plotly y escena 3D"],
        ["App/app.py", "Interfaz, estado, controles y caché"],
        ["report/", "Informe, CSV, figuras y generador reproducible"],
        ["presentation/", "Presentación HTML, ficha y guion por integrante"],
        ["tests/", "28 pruebas de regresión y propiedades"],
    ])
    document.add_heading("Inicio rápido", level=2)
    for command in (
        "cd zombie_poisson_streamlit",
        "python -m venv .venv",
        "python -m pip install -r requirements/requirements.txt",
        "python -m streamlit run App/app.py",
    ):
        paragraph = document.add_paragraph()
        run = paragraph.add_run(command)
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
    document.add_heading("Entregables", level=2)
    add_bullets(document, [
        "Aplicación Streamlit con botón SIMULAR MISIÓN y seis laboratorios.",
        "README con teoría, instalación, resultados y limitaciones.",
        "Informe LaTeX/PDF y CSV regenerables.",
        "Presentación HTML offline y guion LaTeX/PDF dividido entre cinco integrantes.",
        "Ficha DOCX editable y ficha PDF con fuentes versionadas.",
    ])
    document.add_heading("Integrantes", level=2)
    document.add_paragraph(
        "Pablo Daniel Barillas Moreno · Jorge Palacios · "
        "Adrián Ricardo González Muralles · Andrés Rafael Chivalán · "
        "Javier Andrés Chen"
    )
    repository = document.add_paragraph()
    repository.add_run("Repositorio: ").bold = True
    add_hyperlink(repository, REPOSITORY_URL, REPOSITORY_URL)

    document.core_properties.title = "Ficha técnica - Outbreak Stochastic Survival Lab"
    document.core_properties.subject = "Proyecto 1 de Modelación y Simulación"
    document.core_properties.author = "Grupo 1, Sección 30"
    document.save(OUTPUT_PATH)
    print(f"Ficha DOCX regenerada: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
